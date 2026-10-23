"""Resume text extraction and field parsing (pdfplumber, python-docx, regex rules).

Protected attributes (gender, age, photo, religion, address) are intentionally not extracted."""
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from .skills import extract_skills

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(\+?\d[\d\s().\-]{8,16}\d)(?!\d)")
YEARS_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\+?\s*(?:years?|yrs?)(?:\s+of)?(?:\s+(?:professional|work|relevant|industry|total))?"
                      r"(?:\s+of)?\s*(?:experience|exp)", re.I)
MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
RANGE_RE = re.compile(
    rf"(?:(?:{MONTHS})[a-z]*\.?\s+)?((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*"
    rf"(?:(?:(?:{MONTHS})[a-z]*\.?\s+)?((?:19|20)\d{{2}})|(present|current|now|till date))", re.I)

DEGREES = [  # (level rank, label, regex)
    (4, "PhD", r"\bph\.?d\b|doctorate"),
    (3, "Master", r"\bm\.?\s?tech\b|\bm\.?sc\b|\bm\.?s\b(?!\w)|\bmba\b|\bmca\b|\bmaster'?s?\b|\bm\.?e\b(?=\W)"),
    (2, "Bachelor", r"\bb\.?\s?tech\b|\bb\.?e\b(?=\W)|\bb\.?sc\b|\bbca\b|\bb\.?com\b|\bbachelor'?s?\b|\bb\.?a\b(?=\W)"),
    (1, "Diploma", r"\bdiploma\b"),
]
LEVEL_RANK = {"": 0, "diploma": 1, "bachelor": 2, "master": 3, "phd": 4}
SECTION_WORDS = {"resume", "curriculum vitae", "cv", "profile", "summary", "objective", "contact", "skills",
                 "experience", "education", "projects"}


@dataclass
class Parsed:
    name: str = ""
    email: str | None = None
    phone: str = ""
    skills: list[str] = field(default_factory=list)
    experience_years: float = 0.0
    education_level: str = ""
    education: list[str] = field(default_factory=list)
    text: str = ""


def extract_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {ext}")


def _name(lines: list[str], email: str | None) -> str:
    for ln in lines[:8]:
        s = ln.strip()
        if not s or s.lower().strip(":") in SECTION_WORDS or "@" in s or re.search(r"\d", s):
            continue
        words = s.replace(",", " ").split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words):
            return " ".join(w.capitalize() if w.isupper() or w.islower() else w for w in words)
    if email:  # fall back to the mailbox
        return re.sub(r"[._\-]+", " ", email.split("@")[0]).title()
    return ""


def _experience(text: str) -> float:
    explicit = [float(m.group(1)) for m in YEARS_RE.finditer(text)]
    this_year = date.today().year
    spans = []
    for m in RANGE_RE.finditer(text):
        start = int(m.group(1))
        end = int(m.group(2)) if m.group(2) else this_year
        if 1980 <= start <= end <= this_year + 1 and end - start <= 40:
            spans.append((start, end))
    # merge overlapping ranges so concurrent jobs are not double counted
    spans.sort()
    merged: list[list[int]] = []
    for s, e in spans:
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    from_dates = float(sum(e - s for s, e in merged))
    best = max([from_dates] + explicit) if (explicit or merged) else 0.0
    return min(best, 45.0)


def _education(text: str) -> tuple[str, list[str]]:
    levels, lines = [], []
    for ln in text.splitlines():
        for rank, label, rx in DEGREES:
            if re.search(rx, ln, re.I):
                levels.append((rank, label))
                if ln.strip()[:150] not in lines:
                    lines.append(ln.strip()[:150])
                break
    if not levels:
        return "", []
    rank, label = max(levels)
    return label.lower(), lines[:5]


def parse_text(text: str) -> Parsed:
    text = text.replace("\x00", "")
    lines = [l for l in text.splitlines() if l.strip()]
    em = EMAIL_RE.search(text)
    email = em.group(0).lower() if em else None
    ph = PHONE_RE.search(text.replace(em.group(0), "") if em else text)
    level, edu = _education(text)
    return Parsed(
        name=_name(lines, email), email=email,
        phone=re.sub(r"\s+", " ", ph.group(1)).strip() if ph else "",
        skills=extract_skills(text), experience_years=_experience(text),
        education_level=level, education=edu, text=text,
    )


def parse_file(path: str) -> Parsed:
    return parse_text(extract_text(path))
